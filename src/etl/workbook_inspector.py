"""大体积 GTM Excel 的只读 Sheet 识别与按行筛选。"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any, BinaryIO, Iterator

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.models.project import ProjectConfig

HEADER_HINTS = {"country", "project name", "date", "platform", "objective", "impression"}


def select_platform_sheets(sheet_names: list[str]) -> dict[str, str]:
    """从左到右选择 FB、TT 和第一个 GG 数据 Sheet。"""

    selected: dict[str, str] = {}
    for sheet_name in sheet_names:
        normalized = sheet_name.strip().upper()
        if "FB" not in selected and (normalized == "FB" or normalized.startswith("FB ")):
            selected["FB"] = sheet_name
        elif "TT" not in selected and (normalized == "TT" or normalized.startswith("TT ")):
            selected["TT"] = sheet_name
        elif "GG" not in selected and (normalized == "GG" or normalized.startswith("GG ")):
            selected["GG"] = sheet_name
    return selected


def _normalize_header(value: Any) -> str:
    return " ".join(str(value).strip().split()) if value is not None else ""


def _iter_records(sheet: Worksheet, max_header_rows: int = 25) -> Iterator[tuple[int, dict[str, Any]]]:
    headers: list[str] | None = None
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        cleaned = [_normalize_header(value) for value in values]
        if headers is None:
            normalized = {value.lower() for value in cleaned if value}
            if len(normalized & HEADER_HINTS) >= 3:
                headers = cleaned
                continue
            if row_number >= max_header_rows:
                raise ValueError(f"Sheet {sheet.title} 前 {max_header_rows} 行未识别到有效表头")
            continue
        record = {
            header: value
            for header, value in zip(headers, values, strict=False)
            if header
        }
        if any(value not in (None, "") for value in record.values()):
            yield row_number, record


def _matches_project(record: dict[str, Any], project: ProjectConfig, report_date: date | None) -> bool:
    country = str(record.get("Country", "")).strip().upper()
    project_name = str(record.get("Project Name", "")).strip().casefold()
    if country != project.country or project_name != project.project_name.casefold():
        return False
    if report_date is None:
        return True
    parsed = pd.to_datetime(record.get("Date"), errors="coerce")
    return not pd.isna(parsed) and parsed.date() == report_date


def load_gtm_project_data(
    source: str | Path | BinaryIO,
    project: ProjectConfig,
    report_date: date | None = None,
) -> pd.DataFrame:
    """只读取目标平台，并在逐行阶段尽早过滤项目和日期。"""

    workbook = load_workbook(source, read_only=True, data_only=True)
    selected = select_platform_sheets(workbook.sheetnames)
    missing = {"FB", "TT", "GG"} - set(selected)
    if missing:
        workbook.close()
        raise ValueError(f"缺少日报平台 Sheet: {', '.join(sorted(missing))}")

    records: list[dict[str, Any]] = []
    source_name = Path(source).name if isinstance(source, (str, Path)) else "uploaded.xlsx"
    try:
        for platform in ("FB", "TT", "GG"):
            sheet_name = selected[platform]
            sheet = workbook[sheet_name]
            for row_number, record in _iter_records(sheet):
                if _matches_project(record, project, report_date):
                    record["Source File"] = source_name
                    record["Source Sheet"] = sheet_name
                    record["Source Row"] = row_number
                    records.append(record)
    finally:
        workbook.close()
    return pd.DataFrame(records)

